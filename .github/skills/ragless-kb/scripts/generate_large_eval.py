#!/usr/bin/env python3
"""Generate a synthetic large eval corpus (default: 1000 docs, 100 questions).

Writes an isolated workspace (input/, knowledge/nav stubs, eval/questions.json)
so the sample NIST/Northwind corpus is not overwritten. Point the CLI at it with
KB_ROOT.
"""
from __future__ import annotations

import argparse
import json
import random
from pathlib import Path

NOISE = """
Northwind model risk committee reviews production systems annually.
Validation cadence follows the same policy for challenger models unless an
emergency exception is approved. Independent validators sample outcomes,
document residual risk, and file findings with the Model Risk Office.
Shared vocabulary: budget, owner, plant, control, inventory, measurement,
governance, fairness, generative, inventory briefing, production model.
""".strip()


def fact(n: int) -> dict:
    return {
        "n": n,
        "id": f"record-{n:04d}",
        "code": f"qkasset{n:04d}",
        "owner": f"owner{n:04d}name",
        "budget": 100000 + n * 17,
        "site": f"site{n:04d}river",
        "control": f"ctl{n:04d}zeta",
        "title": f"Asset record {n:04d}",
    }


def body(item: dict) -> str:
    return (
        f"# {item['title']}\n\n"
        f"{NOISE}\n\n"
        f"Asset code: {item['code']}\n"
        f"Named owner: {item['owner']}\n"
        f"Operating budget USD: {item['budget']}\n"
        f"Primary site: {item['site']}\n"
        f"Control identifier: {item['control']}\n\n"
        f"This record is the sole source for {item['code']}. Neighboring "
        f"inventory rows reuse the shared Northwind policy language above "
        f"but do not copy this asset code.\n"
    )


def write_txt(path: Path, item: dict) -> None:
    path.write_text(body(item), encoding="utf-8")


def write_md(path: Path, item: dict) -> None:
    path.write_text(body(item), encoding="utf-8")


def write_html(path: Path, item: dict) -> None:
    text = body(item).replace("\n", "<br/>\n")
    path.write_text(
        f"<!doctype html><html><body><h1>{item['title']}</h1><p>{text}</p></body></html>\n",
        encoding="utf-8",
    )


def write_csv(path: Path, item: dict) -> None:
    rows = [
        "field,value",
        f"title,{item['title']}",
        f"asset_code,{item['code']}",
        f"named_owner,{item['owner']}",
        f"operating_budget_usd,{item['budget']}",
        f"primary_site,{item['site']}",
        f"control_identifier,{item['control']}",
        f"notes,{NOISE.replace(',', ' ').replace(chr(10), ' ')}",
    ]
    path.write_text("\n".join(rows) + "\n", encoding="utf-8")


def write_docx(path: Path, item: dict) -> None:
    from docx import Document

    document = Document()
    document.add_heading(item["title"], level=1)
    document.add_paragraph(body(item))
    document.save(path)


def suffix_for(n: int, docx_ok: bool) -> str:
    kind = n % 20
    if kind == 0:
        return ".html"
    if kind == 1:
        return ".csv"
    if kind == 2:
        return ".md"
    if kind == 3 and docx_ok and n <= 200:
        return ".docx"
    return ".txt"


def write_card(path: Path, item: dict, source_name: str) -> None:
    path.write_text(
        (
            f"---\n"
            f"id: {item['id']}\n"
            f"title: {item['title']}\n"
            f"source: knowledge/docs/{item['id']}.md\n"
            f"original: input/{source_name}\n"
            f"authority: reference\n"
            f"status: current\n"
            f"topics: [large-eval]\n"
            f"---\n\n"
            f"# {item['title']}\n\n"
            f"Navigation card for synthetic asset {item['code']}. Not evidence.\n"
        ),
        encoding="utf-8",
    )


def write_index(path: Path, n_docs: int, n_questions: int) -> None:
    path.write_text(
        (
            "# Knowledge base index\n\n"
            f"Synthetic large eval: {n_docs} documents, {n_questions} retrieval questions.\n\n"
            "Topic map: [large eval](topics/large-eval.md)\n\n"
            "Each `record-NNNN` file plants a unique `qkassetNNNN` code. Shared Northwind "
            "policy language is repeated so BM25 must rely on the distinctive token.\n"
        ),
        encoding="utf-8",
    )


def write_topic(path: Path, n_docs: int) -> None:
    path.write_text(
        (
            "# Large eval topic map\n\n"
            f"{n_docs} synthetic asset records. Look up `qkassetNNNN`, `ownerNNNNname`, "
            f"or `ctlNNNNzeta`. Canonical files are `knowledge/docs/record-NNNN.md`.\n"
        ),
        encoding="utf-8",
    )


def build_questions(items: list[dict], n_questions: int, rng: random.Random) -> list[dict]:
    chosen = items[:]
    rng.shuffle(chosen)
    chosen = chosen[:n_questions]
    questions = []
    templates = [
        (
            "code",
            "What is the operating budget for {code}?",
            lambda item: [str(item["budget"])],
        ),
        (
            "owner",
            "Which asset does {owner} own?",
            lambda item: [item["code"]],
        ),
        (
            "control",
            "Which site is covered by control {control}?",
            lambda item: [item["site"]],
        ),
        (
            "site",
            "What control identifier is recorded for {site}?",
            lambda item: [item["control"]],
        ),
        (
            "mixed",
            "In the Northwind inventory, what named owner is listed for {code}?",
            lambda item: [item["owner"]],
        ),
    ]
    for index, item in enumerate(chosen):
        kind, template, gold = templates[index % len(templates)]
        questions.append(
            {
                "id": f"{kind}-{item['id']}",
                "question": template.format(**item),
                "expected_doc_ids": [item["id"]],
                "must_contain": gold(item),
                "gold": f"{item['code']} budget {item['budget']} owner {item['owner']}",
            }
        )
    questions.sort(key=lambda row: row["id"])
    return questions


def generate(root: Path, n_docs: int, n_questions: int, seed: int) -> dict:
    if n_docs < 1 or n_questions < 1:
        raise SystemExit("docs and questions must be >= 1")
    if n_questions > n_docs:
        raise SystemExit("questions cannot exceed docs")
    rng = random.Random(seed)
    inbox = root / "input"
    docs_dir = root / "knowledge" / "docs"
    cards = root / "knowledge" / "cards"
    topics = root / "knowledge" / "topics"
    eval_dir = root / "eval"
    for folder in (inbox, docs_dir, cards, topics, eval_dir):
        folder.mkdir(parents=True, exist_ok=True)
    for leftover in inbox.glob("record-*"):
        leftover.unlink()

    docx_ok = False
    try:
        import docx  # noqa: F401

        docx_ok = True
    except ImportError:
        docx_ok = False

    items = []
    counts: dict[str, int] = {}
    for n in range(1, n_docs + 1):
        item = fact(n)
        items.append(item)
        suffix = suffix_for(n, docx_ok)
        name = f"{item['id']}{suffix}"
        dest = inbox / name
        writers = {
            ".txt": write_txt,
            ".md": write_md,
            ".html": write_html,
            ".csv": write_csv,
            ".docx": write_docx,
        }
        writers[suffix](dest, item)
        write_card(cards / f"{item['id']}.md", item, name)
        counts[suffix] = counts.get(suffix, 0) + 1

    write_index(root / "knowledge" / "INDEX.md", n_docs, n_questions)
    write_topic(topics / "large-eval.md", n_docs)
    questions = build_questions(items, n_questions, rng)
    payload = {
        "k": 5,
        "notes": (
            f"Synthetic large eval ({n_docs} docs, {n_questions} questions, seed={seed}). "
            "Gold is the unique qkassetNNNN token planted in one file."
        ),
        "questions": questions,
    }
    (eval_dir / "questions.json").write_text(json.dumps(payload, indent=2) + "\n", encoding="utf-8")
    summary = {
        "root": str(root),
        "docs": n_docs,
        "questions": n_questions,
        "seed": seed,
        "formats": counts,
        "docx_available": docx_ok,
    }
    print(json.dumps(summary, indent=2))
    return summary


def main() -> None:
    parser = argparse.ArgumentParser(description="Generate a synthetic large retrieval eval")
    parser.add_argument("--root", type=Path, required=True, help="Isolated workspace (set KB_ROOT to this)")
    parser.add_argument("--docs", type=int, default=1000)
    parser.add_argument("--questions", type=int, default=100)
    parser.add_argument("--seed", type=int, default=42)
    args = parser.parse_args()
    generate(args.root.resolve(), args.docs, args.questions, args.seed)


if __name__ == "__main__":
    main()
