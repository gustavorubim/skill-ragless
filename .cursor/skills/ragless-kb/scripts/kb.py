#!/usr/bin/env python3
"""Ragless document knowledge base: ingest, catalog, FTS5 search, validate, eval."""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import os
import re
import sqlite3
import sys
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path

SUPPORTED = {".md", ".markdown", ".txt", ".pdf", ".docx", ".pptx", ".xlsx", ".csv", ".html", ".htm"}
SKIP_NAMES = {"readme.md", ".gitkeep"}
FORMAT_TYPES = {ext.lstrip(".") for ext in SUPPORTED} | {"markdown"}
SEMANTIC_KEYS = (
    "title",
    "document_type",
    "date",
    "version",
    "authority",
    "status",
    "topics",
    "entities",
    "summary",
)
FTS_STOPWORDS = {
    "a", "an", "the", "and", "or", "of", "to", "in", "on", "for", "from", "with",
    "is", "are", "was", "were", "be", "been", "being", "that", "this", "these",
    "those", "it", "its", "as", "at", "by", "into", "about", "what", "which",
    "who", "whom", "how", "why", "when", "where", "do", "does", "did", "can",
    "could", "should", "would", "may", "might", "will", "shall", "me", "my",
    "our", "your", "their", "please", "tell", "give", "show", "find", "list",
    "using", "use", "used", "any", "all", "document", "documents", "kb",
    "often", "must", "receive", "many", "much", "some", "more", "most",
    "does", "also", "than", "each", "has", "have", "including", "within",
    "there", "here", "into", "over", "after", "before", "between",
}
FTS_OPERATORS = {"and", "or", "not", "near"}


@dataclass(frozen=True)
class Paths:
    root: Path
    kb: Path
    inbox: Path
    docs: Path
    cards: Path
    topics: Path
    state: Path
    catalog: Path
    index_md: Path
    db: Path
    manifest: Path
    eval_questions: Path


def get_root() -> Path:
    env = os.environ.get("KB_ROOT")
    if env:
        return Path(env).resolve()
    start = Path(__file__).resolve().parent
    for candidate in [start, *start.parents]:
        if (candidate / "input").is_dir():
            return candidate
    raise RuntimeError(
        "Cannot locate workspace root (no input/ folder). Create input/ at the project root or set KB_ROOT."
    )


def paths() -> Paths:
    root = get_root()
    kb = root / "knowledge"
    state = kb / ".kb"
    return Paths(
        root=root,
        kb=kb,
        inbox=root / "input",
        docs=kb / "docs",
        cards=kb / "cards",
        topics=kb / "topics",
        state=state,
        catalog=kb / "catalog.json",
        index_md=kb / "INDEX.md",
        db=state / "search.sqlite3",
        manifest=state / "manifest.json",
        eval_questions=root / "eval" / "questions.json",
    )


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def slug(value: str) -> str:
    cleaned = re.sub(r"[^a-zA-Z0-9]+", "-", value).strip("-").lower()
    return cleaned[:100] or "document"


def is_dot_path(path: Path, root: Path) -> bool:
    return any(part.startswith(".") for part in path.relative_to(root).parts)


def parse_frontmatter(text: str) -> tuple[dict, str]:
    if not text.startswith("---\n"):
        return {}, text
    end = text.find("\n---\n", 4)
    if end < 0:
        return {}, text
    raw = text[4:end]
    try:
        import yaml

        meta = yaml.safe_load(raw) or {}
    except Exception:
        meta = {}
    if not isinstance(meta, dict):
        meta = {}
    return meta, text[end + 5 :]


def yaml_escape(value: str) -> str:
    return json.dumps(value, ensure_ascii=False)


def frontmatter(meta: dict) -> str:
    lines = ["---"]
    for key, value in meta.items():
        if isinstance(value, list):
            lines.append(f"{key}: [{', '.join(yaml_escape(str(item)) for item in value)}]")
        elif value is None:
            lines.append(f"{key}: null")
        elif isinstance(value, (int, float)) and not isinstance(value, bool):
            lines.append(f"{key}: {value}")
        elif isinstance(value, bool):
            lines.append(f"{key}: {str(value).lower()}")
        else:
            lines.append(f"{key}: {yaml_escape(str(value))}")
    return "\n".join(lines) + "\n---\n\n"


def load_manifest(p: Paths) -> dict:
    if p.manifest.exists():
        try:
            return json.loads(p.manifest.read_text(encoding="utf-8"))
        except Exception:
            pass
    return {"files": {}}


def save_manifest(p: Paths, manifest: dict) -> None:
    p.state.mkdir(parents=True, exist_ok=True)
    p.manifest.write_text(json.dumps(manifest, indent=2), encoding="utf-8")


def is_placeholder(key: str, value) -> bool:
    if value in (None, "", [], {}):
        return True
    if key in ("authority", "status") and value == "unknown":
        return True
    if key == "document_type" and str(value).lower() in FORMAT_TYPES:
        return True
    return False


def merge_semantic_meta(base: dict, existing: dict | None) -> dict:
    if not existing:
        return base
    merged = dict(base)
    for key in SEMANTIC_KEYS:
        value = existing.get(key)
        if is_placeholder(key, value):
            continue
        merged[key] = value
    return merged


def pdf_meta_title(reader) -> str | None:
    info = getattr(reader, "metadata", None)
    if not info:
        return None
    title = getattr(info, "title", None) or (info.get("/Title") if hasattr(info, "get") else None)
    if not title:
        return None
    cleaned = str(title).strip()
    return cleaned or None


def extract(path: Path) -> tuple[str, list[str], dict]:
    ext = path.suffix.lower()
    warnings: list[str] = []
    extra: dict = {}
    if ext in {".md", ".markdown", ".txt"}:
        text = path.read_text(encoding="utf-8", errors="replace")
        if ext != ".txt":
            src_meta, body = parse_frontmatter(text)
            if src_meta:
                extra = {k: v for k, v in src_meta.items() if k in SEMANTIC_KEYS and not is_placeholder(k, v)}
                return body, warnings, extra
        return text, warnings, extra
    if ext == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise RuntimeError("Install pypdf from requirements.txt") from exc
        reader = PdfReader(str(path))
        title = pdf_meta_title(reader)
        if title:
            extra["title"] = title
        parts = []
        for index, page in enumerate(reader.pages, 1):
            parts += [f"\n<!-- source-page: {index} -->\n", page.extract_text() or ""]
        text = "\n".join(parts)
        if len(text.strip()) < 200:
            warnings.append("PDF extraction produced very little text; source may be scanned/image-only.")
        return text, warnings, extra
    if ext == ".docx":
        try:
            from docx import Document
        except ImportError as exc:
            raise RuntimeError("Install python-docx from requirements.txt") from exc
        document = Document(str(path))
        out = []
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue
            style = (paragraph.style.name or "").lower() if paragraph.style else ""
            if style.startswith("heading"):
                match = re.search(r"(\d+)", style)
                level = min(6, int(match.group(1))) if match else 2
                out.append("#" * level + " " + text)
            else:
                out.append(text)
        for table_index, table in enumerate(document.tables, 1):
            out.append(f"\n### Table {table_index}")
            rows = [[cell.text.replace("\n", " ").strip() for cell in row.cells] for row in table.rows]
            if rows:
                out.append("| " + " | ".join(rows[0]) + " |")
                out.append("| " + " | ".join(["---"] * len(rows[0])) + " |")
                out += ["| " + " | ".join(row) + " |" for row in rows[1:]]
        return "\n\n".join(out), warnings, extra
    if ext == ".pptx":
        try:
            from pptx import Presentation
        except ImportError as exc:
            raise RuntimeError("Install python-pptx from requirements.txt") from exc
        presentation = Presentation(str(path))
        out = []
        for index, slide in enumerate(presentation.slides, 1):
            out.append(f"\n<!-- source-slide: {index} -->\n## Slide {index}")
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    out.append(shape.text.strip())
        return "\n\n".join(out), warnings, extra
    if ext == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError("Install openpyxl from requirements.txt") from exc
        workbook = load_workbook(str(path), data_only=False, read_only=True)
        out = []
        for sheet in workbook.worksheets:
            out.append(f"\n## Sheet: {sheet.title}")
            rows = list(sheet.iter_rows(values_only=True))
            for row in rows[:10000]:
                out.append("| " + " | ".join("" if value is None else str(value).replace("\n", " ") for value in row) + " |")
            if len(rows) > 10000:
                warnings.append(f"Sheet {sheet.title} truncated at 10,000 rows in Markdown conversion.")
        return "\n".join(out), warnings, extra
    if ext == ".csv":
        out = []
        with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as handle:
            for index, row in enumerate(csv.reader(handle)):
                if index >= 10000:
                    warnings.append("CSV truncated at 10,000 rows in Markdown conversion.")
                    break
                out.append("| " + " | ".join(item.replace("\n", " ") for item in row) + " |")
        return "\n".join(out), warnings, extra
    if ext in {".html", ".htm"}:
        try:
            from bs4 import BeautifulSoup
        except ImportError as exc:
            raise RuntimeError("Install beautifulsoup4 from requirements.txt") from exc
        soup = BeautifulSoup(path.read_text(encoding="utf-8", errors="replace"), "html.parser")
        for node in soup(["script", "style"]):
            node.decompose()
        return soup.get_text("\n", strip=True), warnings, extra
    raise RuntimeError(f"Unsupported extension: {ext}")


def default_meta(doc_id: str, source: Path, rel: str, digest: str, warns: list[str]) -> dict:
    title = source.stem.replace("_", " ").replace("-", " ").strip()
    return {
        "id": doc_id,
        "title": title,
        "source_file": rel,
        "source_sha256": digest,
        "document_type": source.suffix.lower().lstrip("."),
        "date": None,
        "version": None,
        "authority": "unknown",
        "status": "unknown",
        "topics": [],
        "entities": [],
        "summary": "",
        "extraction_warnings": warns,
    }


def ensure_dirs(p: Paths) -> None:
    for folder in (p.inbox, p.docs, p.cards, p.topics, p.state):
        folder.mkdir(parents=True, exist_ok=True)


def ingest() -> int:
    p = paths()
    ensure_dirs(p)
    manifest = load_manifest(p)
    files = manifest.setdefault("files", {})
    stats = {"added": 0, "updated": 0, "unchanged": 0, "failed": 0, "unsupported": 0, "orphaned": 0}
    seen: set[str] = set()

    for source in sorted(item for item in p.inbox.rglob("*") if item.is_file() and not is_dot_path(item, p.inbox)):
        if source.name.lower() in SKIP_NAMES:
            continue
        rel = source.relative_to(p.root).as_posix()
        seen.add(rel)
        ext = source.suffix.lower()
        if ext not in SUPPORTED:
            print(f"UNSUPPORTED {rel}")
            stats["unsupported"] += 1
            continue
        digest = sha256(source)
        prior = files.get(rel)
        base = slug(source.stem)
        doc_id = (prior or {}).get("id") or base
        dest = p.docs / f"{doc_id}.md"
        if prior and prior.get("sha256") == digest and dest.exists():
            stats["unchanged"] += 1
            continue
        if dest.exists() and (not prior or prior.get("id") != doc_id):
            doc_id = f"{base}-{digest[:8]}"
            dest = p.docs / f"{doc_id}.md"
        try:
            text, warns, source_meta = extract(source)
            meta = default_meta(doc_id, source, rel, digest, warns)
            meta = merge_semantic_meta(meta, source_meta)
            if dest.exists():
                existing, _ = parse_frontmatter(dest.read_text(encoding="utf-8", errors="replace"))
                meta = merge_semantic_meta(meta, existing)
            title = meta.get("title") or source.stem
            body = frontmatter(meta) + f"# {title}\n\n{text.strip()}\n"
            dest.write_text(body, encoding="utf-8")
            files[rel] = {
                "sha256": digest,
                "id": doc_id,
                "doc": dest.relative_to(p.root).as_posix(),
                "source": rel,
                "updated_at": datetime.now(timezone.utc).isoformat(),
                "orphaned": False,
            }
            stats["updated" if prior else "added"] += 1
            print(f"OK {rel} -> {dest.relative_to(p.root)}")
        except Exception as exc:
            stats["failed"] += 1
            print(f"FAILED {rel}: {exc}", file=sys.stderr)

    for rel, rec in files.items():
        if rel in seen:
            rec["orphaned"] = False
            continue
        rec["orphaned"] = True
        stats["orphaned"] += 1
        print(f"ORPHAN source missing: {rel} (canonical {rec.get('doc')} retained)")

    manifest["last_ingest"] = datetime.now(timezone.utc).isoformat()
    save_manifest(p, manifest)
    print(json.dumps(stats, indent=2))
    return 1 if stats["failed"] else 0


def load_docs(p: Paths) -> list[dict]:
    docs = []
    for doc_path in sorted(p.docs.glob("*.md")):
        text = doc_path.read_text(encoding="utf-8", errors="replace")
        meta, body = parse_frontmatter(text)
        meta = dict(meta)
        meta.setdefault("id", doc_path.stem)
        meta["path"] = doc_path.relative_to(p.root).as_posix()
        meta["headings"] = re.findall(r"^#{1,6}\s+(.+)$", body, re.M)
        docs.append(meta)
    return docs


def rebuild() -> None:
    p = paths()
    ensure_dirs(p)
    docs = load_docs(p)
    p.catalog.write_text(json.dumps(docs, indent=2, ensure_ascii=False), encoding="utf-8")
    if p.db.exists():
        p.db.unlink()
    con = sqlite3.connect(p.db)
    con.execute(
        "CREATE VIRTUAL TABLE docs USING fts5("
        "id UNINDEXED, path UNINDEXED, source_file UNINDEXED, title, summary, topics, body, tokenize='unicode61')"
    )
    for meta in docs:
        doc_path = p.root / meta["path"]
        _, body = parse_frontmatter(doc_path.read_text(encoding="utf-8", errors="replace"))
        topics = meta.get("topics") or []
        topic_text = topics if isinstance(topics, str) else " ".join(str(item) for item in topics)
        con.execute(
            "INSERT INTO docs(id, path, source_file, title, summary, topics, body) VALUES (?, ?, ?, ?, ?, ?, ?)",
            (
                meta.get("id", ""),
                meta["path"],
                meta.get("source_file", "") or "",
                meta.get("title", ""),
                meta.get("summary", "") or "",
                topic_text,
                body,
            ),
        )
    con.commit()
    con.close()
    print(f"Rebuilt catalog ({len(docs)} docs) and {p.db.relative_to(p.root)}")


def kept_terms(raw: str) -> list[str]:
    terms = re.findall(r"[A-Za-z0-9_]+", raw or "")
    kept = []
    for term in terms:
        lower = term.lower()
        if lower in FTS_STOPWORDS or lower in FTS_OPERATORS or len(lower) < 2:
            continue
        if lower not in kept:
            kept.append(lower)
    return kept


def fts_query(raw: str) -> str:
    query = (raw or "").strip()
    if not query:
        return query
    kept = kept_terms(query)
    if not kept:
        return '"' + query.replace('"', " ") + '"'
    ranked = sorted(kept, key=lambda term: (-len(term), term))[:10]
    return " OR ".join(ranked)


def fts_query_and(raw: str) -> str:
    kept = kept_terms(raw)
    if not kept:
        return fts_query(raw)
    return " AND ".join(kept[:8])


def search_rows(query: str, limit: int) -> list[dict]:
    p = paths()
    if not p.db.exists():
        rebuild()
    con = sqlite3.connect(p.db)
    sql = (
        "SELECT id, path, source_file, title, snippet(docs, 6, '[', ']', ' … ', 18), "
        "bm25(docs, 0, 0, 0, 4.0, 1.5, 2.0, 1.0) score "
        "FROM docs WHERE docs MATCH ? ORDER BY score LIMIT ?"
    )

    def run(match: str) -> list[tuple]:
        try:
            return con.execute(sql, (match, limit)).fetchall()
        except sqlite3.OperationalError:
            quoted = '"' + query.replace('"', " ") + '"'
            return con.execute(sql, (quoted, limit)).fetchall()

    rows = run(fts_query(query))
    if not rows:
        rows = run(fts_query_and(query))
    con.close()
    results = []
    for doc_id, path, source_file, title, snippet, score in rows:
        results.append(
            {
                "id": doc_id,
                "path": path,
                "source_file": source_file,
                "title": title or doc_id,
                "snippet": snippet,
                "score": score,
            }
        )
    return results


def search(query: str, limit: int, as_json: bool) -> None:
    rows = search_rows(query, limit)
    if as_json:
        print(json.dumps(rows, indent=2, ensure_ascii=False))
        return
    if not rows:
        print("No matches.")
        return
    for index, row in enumerate(rows, 1):
        print(
            f"{index}. {row['title']}\n"
            f"   {row['path']}\n"
            f"   original: {row['source_file']}\n"
            f"   score={row['score']:.4f}\n"
            f"   {row['snippet']}\n"
        )


def status() -> None:
    p = paths()
    manifest = load_manifest(p)
    inbox = [item for item in p.inbox.rglob("*") if item.is_file() and not is_dot_path(item, p.inbox)]
    print(
        json.dumps(
            {
                "input_files": len(inbox),
                "canonical_docs": len(list(p.docs.glob("*.md"))),
                "cards": len(list(p.cards.glob("*.md"))),
                "topic_maps": len(list(p.topics.glob("*.md"))),
                "index_exists": p.index_md.exists(),
                "catalog_exists": p.catalog.exists(),
                "search_index_exists": p.db.exists(),
                "orphaned_sources": sum(1 for rec in manifest.get("files", {}).values() if rec.get("orphaned")),
                "last_ingest": manifest.get("last_ingest"),
                "python": sys.executable,
            },
            indent=2,
        )
    )


def validate() -> int:
    p = paths()
    problems = []
    manifest = load_manifest(p).get("files", {})
    for rel, rec in manifest.items():
        source = p.root / rel
        dest = p.root / rec.get("doc", "")
        if rec.get("orphaned") or not source.exists():
            problems.append(f"Orphaned source (file missing, canonical retained): {rel}")
        if not dest.exists():
            problems.append(f"Missing canonical doc for {rel}: {dest}")
    for doc_path in p.docs.glob("*.md"):
        meta, body = parse_frontmatter(doc_path.read_text(encoding="utf-8", errors="replace"))
        if not meta.get("source_file"):
            problems.append(f"{doc_path.relative_to(p.root)}: missing source_file")
        if not meta.get("source_sha256"):
            problems.append(f"{doc_path.relative_to(p.root)}: missing source_sha256")
        if len(body.strip()) < 100:
            problems.append(f"{doc_path.relative_to(p.root)}: very short extracted body ({len(body.strip())} chars)")
        card = p.cards / f"{meta.get('id', doc_path.stem)}.md"
        if not card.exists():
            problems.append(f"{doc_path.relative_to(p.root)}: missing card {card.relative_to(p.root)}")
        source_file = meta.get("source_file")
        if source_file and not (p.root / source_file).exists():
            problems.append(f"{doc_path.relative_to(p.root)}: source_file does not exist: {source_file}")
    if not p.catalog.exists():
        problems.append("Missing knowledge/catalog.json; run rebuild")
    if not p.index_md.exists():
        problems.append("Missing knowledge/INDEX.md; the preprocess skill should create the corpus map")
    print("VALID" if not problems else "VALIDATION WARNINGS/ERRORS")
    for item in problems:
        print(f"- {item}")
    return 1 if problems else 0


def load_questions(path: Path) -> list[dict]:
    payload = json.loads(path.read_text(encoding="utf-8"))
    if isinstance(payload, list):
        return payload
    return payload.get("questions") or []


def eval_retrieval(k: int, as_json: bool) -> int:
    p = paths()
    if not p.eval_questions.exists():
        print(f"Missing {p.eval_questions.relative_to(p.root)}", file=sys.stderr)
        return 1
    if not p.db.exists():
        rebuild()
    questions = load_questions(p.eval_questions)
    results = []
    hits_at_k = 0
    snippet_hits = 0
    snippet_total = 0
    for item in questions:
        qid = item.get("id") or item.get("question")
        question = item["question"]
        expected = [str(x) for x in item.get("expected_doc_ids") or []]
        rows = search_rows(question, k)
        found_ids = [row["id"] for row in rows]
        if expected:
            hit = any(doc_id in found_ids for doc_id in expected)
            if hit:
                hits_at_k += 1
        else:
            hit = None
        must = item.get("must_contain") or []
        snippet_ok = None
        if must:
            snippet_total += 1
            haystacks = []
            for doc_id in expected or found_ids[:1]:
                doc_path = p.docs / f"{doc_id}.md"
                if doc_path.exists():
                    haystacks.append(doc_path.read_text(encoding="utf-8", errors="replace"))
            blob = "\n".join(haystacks).lower()
            snippet_ok = all(str(token).lower() in blob for token in must)
            if snippet_ok:
                snippet_hits += 1
        results.append(
            {
                "id": qid,
                "question": question,
                "expected_doc_ids": expected,
                "hit_at_k": hit,
                "found_ids": found_ids,
                "snippet_ok": snippet_ok,
                "top_hit": rows[0] if rows else None,
            }
        )
    n = len([q for q in questions if q.get("expected_doc_ids")])
    report = {
        "k": k,
        "questions": len(questions),
        "retrieval_n": n,
        "hit_at_k": hits_at_k,
        "hit_at_k_rate": (hits_at_k / n) if n else None,
        "snippet_n": snippet_total,
        "snippet_hits": snippet_hits,
        "snippet_rate": (snippet_hits / snippet_total) if snippet_total else None,
        "results": results,
    }
    if as_json:
        print(json.dumps(report, indent=2, ensure_ascii=False))
    else:
        rate = report["hit_at_k_rate"]
        snippet_rate = report["snippet_rate"]
        print(f"Retrieval hit@{k}: {hits_at_k}/{n}" + (f" ({rate:.0%})" if rate is not None else ""))
        if snippet_total:
            print(f"Gold snippet coverage: {snippet_hits}/{snippet_total}" + (f" ({snippet_rate:.0%})" if snippet_rate is not None else ""))
        for item in results:
            mark = "PASS" if item["hit_at_k"] else "FAIL"
            if item["hit_at_k"] is None:
                mark = "SKIP"
            extra = ""
            if item["snippet_ok"] is False:
                extra = " SNIPPET-MISS"
            elif item["snippet_ok"] is True:
                extra = " SNIPPET-OK"
            print(f"- {mark}{extra} {item['id']}: {item['found_ids'][:k]}")
    return 0 if n and hits_at_k == n and (not snippet_total or snippet_hits == snippet_total) else 1


def preferred_python() -> None:
    p = paths()
    win = p.root / ".venv" / "Scripts" / "python.exe"
    unix = p.root / ".venv" / "bin" / "python"
    if win.exists():
        print(str(win))
    elif unix.exists():
        print(str(unix))
    else:
        print(sys.executable)


def main() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        try:
            sys.stdout.reconfigure(errors="replace")
            sys.stderr.reconfigure(errors="replace")
        except Exception:
            pass
    parser = argparse.ArgumentParser(description="Ragless document knowledge base")
    sub = parser.add_subparsers(dest="cmd", required=True)
    sub.add_parser("ingest", help="Convert input/ files into canonical Markdown")
    sub.add_parser("preprocess", help="Alias for ingest")
    sub.add_parser("rebuild", help="Regenerate catalog.json and the FTS5 index")
    sub.add_parser("status", help="Print corpus counts and last ingest time")
    sub.add_parser("validate", help="Check coverage, provenance, cards, index, and orphans")
    sub.add_parser("python", help="Print .venv interpreter if present (prefer uv run python)")
    sub.add_parser("skill-path", help="Print this kb.py path")
    search_parser = sub.add_parser("search", help="BM25/FTS5 search over canonical docs")
    search_parser.add_argument("query")
    search_parser.add_argument("--limit", type=int, default=10)
    search_parser.add_argument("--json", action="store_true")
    eval_parser = sub.add_parser("eval", help="Run retrieval eval against eval/questions.json")
    eval_parser.add_argument("--k", type=int, default=5)
    eval_parser.add_argument("--json", action="store_true")
    args = parser.parse_args()
    if args.cmd in {"ingest", "preprocess"}:
        raise SystemExit(ingest())
    if args.cmd == "rebuild":
        rebuild()
    elif args.cmd == "search":
        search(args.query, args.limit, args.json)
    elif args.cmd == "status":
        status()
    elif args.cmd == "python":
        preferred_python()
    elif args.cmd == "skill-path":
        print(str(Path(__file__).resolve()))
    elif args.cmd == "validate":
        raise SystemExit(validate())
    elif args.cmd == "eval":
        raise SystemExit(eval_retrieval(args.k, args.json))


if __name__ == "__main__":
    main()
