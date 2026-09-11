#!/usr/bin/env python3
"""Create small office/text fixtures with known gold facts for eval."""
from __future__ import annotations

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from kb import get_root

ROOT = get_root()
INPUT = ROOT / "input"


def write_txt() -> None:
    text = """# Northwind Model Risk FAQ

Last updated: 2026-06-15
Owner: Model Risk Management Office

## Validation cadence

Challenger models follow the same annual independent validation cadence as production models
unless an emergency exception is approved by the Model Risk Committee (MRC).

## Emergency exceptions

Emergency exceptions expire after ninety (90) days. A model may not remain in production
on an expired exception.

## Contact

Questions: model-risk@northwind.example
"""
    (INPUT / "faq-model-risk.txt").write_text(text, encoding="utf-8")


def write_csv() -> None:
    text = """model_id,name,owner,risk_tier,status,last_validation
NW-001,Credit-Decision-v4,Retail Credit,high,production,2025-11-02
NW-014,Fraud-Graph-v2,Payments,high,production,2026-01-18
NW-022,Marketing-Propensity-v9,Growth,medium,production,2025-08-30
NW-031,HR-Screening-Assist,People Analytics,high,validation-queue,never
NW-044,GenAI-Contract-Summarizer,Legal Ops,high,production,2026-03-01
"""
    (INPUT / "model-inventory.csv").write_text(text, encoding="utf-8")


def write_docx() -> None:
    from docx import Document

    doc = Document()
    doc.add_heading("Northwind AI Governance Policy", level=1)
    doc.add_paragraph("Document ID: NW-AIGP-3.2")
    doc.add_paragraph("Version: 3.2")
    doc.add_paragraph("Effective date: 2026-04-01")
    doc.add_paragraph("Classification: Internal policy")
    doc.add_heading("Purpose", level=2)
    doc.add_paragraph(
        "This policy governs how Northwind Bank develops, validates, and operates artificial "
        "intelligence and machine learning systems. It is an internal control standard. "
        "Where this policy is quieter than applicable law or a formal regulator framework, "
        "the stricter requirement wins."
    )
    doc.add_heading("Independent validation", level=2)
    doc.add_paragraph(
        "Every model in production must receive independent validation at least annually. "
        "The Model Risk Committee (MRC) is the approval authority for production use, "
        "material change, and retirement."
    )
    doc.add_heading("Emergency exceptions", level=2)
    doc.add_paragraph(
        "The MRC may approve an emergency production exception. Emergency exceptions expire "
        "after ninety (90) days and cannot be silently renewed. A written close-out or a "
        "full validation is required before expiry."
    )
    doc.add_heading("MEASURE function exemption", level=2)
    doc.add_paragraph(
        "AI systems with forecast annual spend under USD 250,000 are exempt from the MEASURE "
        "function of the NIST AI Risk Management Framework. GOVERN, MAP, and MANAGE still apply. "
        "This exemption is a Northwind cost-control choice and is not stated in NIST AI RMF 1.0."
    )
    doc.add_heading("Generative AI", level=2)
    doc.add_paragraph(
        "Generative AI systems that draft customer-facing content require human review before "
        "send. The GenAI-Contract-Summarizer (NW-044) is in scope."
    )
    path = INPUT / "northwind-ai-governance-policy.docx"
    doc.save(path)


def write_pptx() -> None:
    from pptx import Presentation

    prs = Presentation()
    layout = prs.slide_layouts[1]

    s1 = prs.slides.add_slide(layout)
    s1.shapes.title.text = "Q3 2026 Model Inventory Briefing"
    s1.placeholders[1].text = (
        "Northwind Model Risk Committee\n"
        "47 production models\n"
        "12 models in the validation queue\n"
        "Briefing date: 2026-09-08"
    )

    s2 = prs.slides.add_slide(layout)
    s2.shapes.title.text = "Highest-risk production model"
    s2.placeholders[1].text = (
        "Credit-Decision-v4 (NW-001)\n"
        "Owner: Retail Credit\n"
        "Last independent validation: 2025-11-02\n"
        "Next validation due: 2026-11-02\n"
        "Status: production"
    )

    s3 = prs.slides.add_slide(layout)
    s3.shapes.title.text = "NIST AI RMF ownership map"
    s3.placeholders[1].text = (
        "GOVERN owner: Chief Risk Officer\n"
        "MAP owner: Model Risk Management Office\n"
        "MEASURE owner: Independent Validation\n"
        "MANAGE owner: First-line model owners"
    )

    path = INPUT / "q3-model-inventory.pptx"
    prs.save(path)


def main() -> None:
    INPUT.mkdir(parents=True, exist_ok=True)
    write_txt()
    write_csv()
    write_docx()
    write_pptx()
    print(f"Wrote sample office/text files into {INPUT}")


if __name__ == "__main__":
    main()
